import pandas as pd
import sqlite3
import os
import warnings
from typing import Optional

warnings.filterwarnings("ignore", category=FutureWarning)
warnings.filterwarnings("ignore", category=DeprecationWarning)

from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .database import Database
from .classifiers import classify_essential


class Reporter:
    def __init__(self, db: Database, output_dir: str = "data/reports"):
        self.db = db
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)


    def _apply_header_style(self, ws, min_row, min_col, max_col, title_bg="1F4E78", font_size=11):
        header_fill = PatternFill(start_color=title_bg, end_color=title_bg, fill_type="solid")
        header_font = Font(name="Calibri", size=font_size, bold=True, color="FFFFFF")
        border_bottom = Border(bottom=Side(style='medium', color="000000"))
        
        for col in range(min_col, max_col + 1):
            cell = ws.cell(row=min_row, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border_bottom

    def _format_table(self, ws, min_row, max_row, col_formats):
        """Aplica formatação monetária, percentual e bordas leves."""
        thin_border = Border(
            left=Side(style='thin', color="E0E0E0"),
            right=Side(style='thin', color="E0E0E0"),
            top=Side(style='thin', color="E0E0E0"),
            bottom=Side(style='thin', color="E0E0E0")
        )
        for row in range(min_row, max_row + 1):
            for col_idx, fmt in col_formats.items():
                cell = ws.cell(row=row, column=col_idx)
                if fmt == 'currency':
                    cell.number_format = '"€"#,##0.00'
                    cell.alignment = Alignment(horizontal="right")
                elif fmt == 'percent':
                    cell.number_format = '0.0%'
                    cell.alignment = Alignment(horizontal="right")
                elif fmt == 'integer':
                    cell.number_format = '#,##0'
                    cell.alignment = Alignment(horizontal="right")
                elif fmt == 'decimal':
                    cell.number_format = '#,##0.00'
                    cell.alignment = Alignment(horizontal="right")
                elif fmt == 'center':
                    cell.alignment = Alignment(horizontal="center")
                
                cell.border = thin_border

    def _auto_adjust_columns(self, ws, max_width=45):
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), max_width)

    def generate_excel_report(self, start_date: Optional[str] = None, end_date: Optional[str] = None, filename: Optional[str] = None) -> str:
        from datetime import datetime
        import shutil

        # Gerar nome automático baseado nas datas e timestamp para cada execução ser única
        if not filename:
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")
            if start_date and end_date:
                filename = f"relatorio_{start_date}_a_{end_date}_{ts}.xlsx"
            elif start_date:
                filename = f"relatorio_desde_{start_date}_{ts}.xlsx"
            elif end_date:
                filename = f"relatorio_ate_{end_date}_{ts}.xlsx"
            else:
                filename = f"relatorio_global_{ts}.xlsx"

        output_path = os.path.join(self.output_dir, filename)
        latest_path = os.path.join(self.output_dir, "relatorio_faturas.xlsx")

        query = """
            SELECT 
                i.invoice_date as Data,
                i.store as Loja,
                i.invoice_number as Fatura,
                l.category as Categoria,
                l.description as Produto,
                l.quantity as Quantidade,
                l.unit_price as Preco_Unitario,
                l.total_price as Total,
                l.discount as Desconto
            FROM line_items l
            JOIN invoices i ON l.invoice_id = i.id
            WHERE 1=1
        """
        params = []
        if start_date:
            query += " AND i.invoice_date >= ?"
            params.append(start_date)
        if end_date:
            query += " AND i.invoice_date <= ?"
            params.append(end_date)
        query += " ORDER BY i.invoice_date DESC"

        df = pd.read_sql_query(query, self.db.conn, params=params)
        if df.empty:
            print("Nenhum dado encontrado para o período especificado.")
            return ""

        df['Data'] = pd.to_datetime(df['Data'])
        df['Mês'] = df['Data'].dt.to_period('M').astype(str)
        df['Categoria'] = df['Categoria'].fillna('Outros')
        for col in ['Quantidade', 'Total']:
            df[col] = pd.to_numeric(df[col], errors='coerce').fillna(0)

        total_global = float(df['Total'].sum())

        # Classificação Cesta Básica
        df['Subgrupo_Cesta'] = df.apply(classify_essential, axis=1)
        df['Eh_Cesta_Basica'] = df['Subgrupo_Cesta'].notnull()

        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            # ==============================================================
            # 1. ABA VISÃO GERAL (Dashboard Executivo)
            # ==============================================================
            resumo_mes = df.groupby('Mês')['Total'].sum().reset_index()
            resumo_mes.columns = ['Mês', 'Total_Gasto']

            top_cats = df.groupby('Categoria')['Total'].sum().sort_values(ascending=False).head(10).reset_index()
            top_cats['%_do_Orçamento'] = (top_cats['Total'] / total_global).round(3)
            top_cats.columns = ['Categoria', 'Gasto_Total', '%_do_Orçamento']

            # Resumo Essenciais vs Outros
            total_essencial = float(df[df['Eh_Cesta_Basica']]['Total'].sum())
            total_outros = total_global - total_essencial
            resumo_split = pd.DataFrame([
                {'Tipo': 'Cesta Básica & Essenciais', 'Gasto': total_essencial, '%': total_essencial / total_global},
                {'Tipo': 'Outros Gastos / Supérfluos', 'Gasto': total_outros, '%': total_outros / total_global}
            ])

            ws_dashboard = writer.book.create_sheet(title="1. Visão Geral")
            ws_dashboard.views.sheetView[0].showGridLines = True

            # Título do Dashboard
            ws_dashboard['B2'] = "DASHBOARD FINANCEIRO DE SUPERMERCADO"
            ws_dashboard['B2'].font = Font(name="Calibri", size=16, bold=True, color="1F4E78")
            ws_dashboard['B3'] = f"Gasto Global Analisado: € {total_global:,.2f}  |  Período: {df['Mês'].min()} até {df['Mês'].max()}"
            ws_dashboard['B3'].font = Font(name="Calibri", size=11, italic=True, color="595959")

            # 1.1 Tabela de Meses
            ws_dashboard['B5'] = "Gasto Mensal"
            ws_dashboard['B5'].font = Font(name="Calibri", size=12, bold=True)
            ws_dashboard.append([])  # blank row handled by coordinates
            start_r = 6
            ws_dashboard.cell(row=start_r, column=2, value="Mês")
            ws_dashboard.cell(row=start_r, column=3, value="Total Gasto")
            self._apply_header_style(ws_dashboard, min_row=start_r, min_col=2, max_col=3)
            curr_r = start_r + 1
            for _, r in resumo_mes.iterrows():
                ws_dashboard.cell(row=curr_r, column=2, value=str(r['Mês']))
                ws_dashboard.cell(row=curr_r, column=3, value=float(r['Total_Gasto']))
                curr_r += 1
            # Total Row
            ws_dashboard.cell(row=curr_r, column=2, value="TOTAL GERAL").font = Font(bold=True)
            ws_dashboard.cell(row=curr_r, column=3, value=total_global).font = Font(bold=True)
            self._format_table(ws_dashboard, min_row=start_r+1, max_row=curr_r, col_formats={2: 'center', 3: 'currency'})
            mes_end_r = curr_r

            # 1.2 Tabela Cesta Básica vs Outros
            split_r = mes_end_r + 3
            ws_dashboard.cell(row=split_r, column=2, value="Divisão: Essenciais vs Outros").font = Font(name="Calibri", size=12, bold=True)
            split_header_r = split_r + 1
            ws_dashboard.cell(row=split_header_r, column=2, value="Tipo de Despesa")
            ws_dashboard.cell(row=split_header_r, column=3, value="Gasto Total")
            ws_dashboard.cell(row=split_header_r, column=4, value="% Orçamento")
            self._apply_header_style(ws_dashboard, min_row=split_header_r, min_col=2, max_col=4, title_bg="2E75B6")
            curr_r = split_header_r + 1
            for _, r in resumo_split.iterrows():
                ws_dashboard.cell(row=curr_r, column=2, value=str(r['Tipo']))
                ws_dashboard.cell(row=curr_r, column=3, value=float(r['Gasto']))
                ws_dashboard.cell(row=curr_r, column=4, value=float(r['%']))
                curr_r += 1
            self._format_table(ws_dashboard, min_row=split_header_r+1, max_row=curr_r-1, col_formats={2: 'left', 3: 'currency', 4: 'percent'})
            split_end_r = curr_r - 1

            # 1.3 Tabela Top Categorias
            cat_header_r = 6
            ws_dashboard.cell(row=cat_header_r-1, column=6, value="Top 10 Categorias de Gasto").font = Font(name="Calibri", size=12, bold=True)
            ws_dashboard.cell(row=cat_header_r, column=6, value="Categoria")
            ws_dashboard.cell(row=cat_header_r, column=7, value="Total Gasto")
            ws_dashboard.cell(row=cat_header_r, column=8, value="% Orçamento")
            self._apply_header_style(ws_dashboard, min_row=cat_header_r, min_col=6, max_col=8)
            curr_r = cat_header_r + 1
            for _, r in top_cats.iterrows():
                ws_dashboard.cell(row=curr_r, column=6, value=str(r['Categoria']))
                ws_dashboard.cell(row=curr_r, column=7, value=float(r['Gasto_Total']))
                ws_dashboard.cell(row=curr_r, column=8, value=float(r['%_do_Orçamento']))
                curr_r += 1
            self._format_table(ws_dashboard, min_row=cat_header_r+1, max_row=curr_r-1, col_formats={6: 'left', 7: 'currency', 8: 'percent'})
            cat_end_r = curr_r - 1

            # ==============================================================
            # GRÁFICOS NO DASHBOARD
            # ==============================================================
            # Gráfico 1: Evolução Mensal (BarChart Colunas)
            chart_mes = BarChart()
            chart_mes.type = "col"
            chart_mes.style = 10
            chart_mes.title = "Evolução Mensal dos Gastos (€)"
            chart_mes.y_axis.title = "Gasto (€)"
            chart_mes.x_axis.title = "Mês"
            chart_mes.width = 15
            chart_mes.height = 9
            chart_mes.legend = None
            data_mes = Reference(ws_dashboard, min_col=3, min_row=start_r, max_row=mes_end_r-1)
            cats_mes = Reference(ws_dashboard, min_col=2, min_row=start_r+1, max_row=mes_end_r-1)
            chart_mes.add_data(data_mes, titles_from_data=True)
            chart_mes.set_categories(cats_mes)
            ws_dashboard.add_chart(chart_mes, "J5")

            # Gráfico 2: Cesta Básica vs Outros (PieChart)
            chart_split = PieChart()
            chart_split.title = "Cesta Básica vs Outros Gastos"
            chart_split.width = 14
            chart_split.height = 9
            data_split = Reference(ws_dashboard, min_col=3, min_row=split_header_r, max_row=split_end_r)
            cats_split = Reference(ws_dashboard, min_col=2, min_row=split_header_r+1, max_row=split_end_r)
            chart_split.add_data(data_split, titles_from_data=True)
            chart_split.set_categories(cats_split)
            ws_dashboard.add_chart(chart_split, "J21")

            self._auto_adjust_columns(ws_dashboard)

            # ==============================================================
            # 2. ABA CESTA BÁSICA & ESSENCIAIS (Requisito Explícito)
            # ==============================================================
            df_cesta = df[df['Eh_Cesta_Basica']].copy()
            ws_cesta = writer.book.create_sheet(title="2. Cesta Básica")
            ws_cesta.views.sheetView[0].showGridLines = True

            ws_cesta['B2'] = "RESUMO DA CESTA BÁSICA & GASTOS ESSENCIAIS"
            ws_cesta['B2'].font = Font(name="Calibri", size=16, bold=True, color="1F4E78")
            ws_cesta['B3'] = f"Total Cesta Básica: € {total_essencial:,.2f}  |  Representa {total_essencial/total_global*100:.1f}% do seu carrinho de supermercado"
            ws_cesta['B3'].font = Font(name="Calibri", size=11, bold=True, color="2E75B6")

            # Tabela Resumo por Subgrupo da Cesta Básica
            cesta_subgrupo = df_cesta.groupby(['Subgrupo_Cesta', 'Mês'])['Total'].sum().reset_index()
            cesta_pivot = cesta_subgrupo.pivot(index='Subgrupo_Cesta', columns='Mês', values='Total').fillna(0)
            cesta_pivot['Total_Acumulado'] = cesta_pivot.sum(axis=1)
            cesta_pivot['%_da_Cesta'] = (cesta_pivot['Total_Acumulado'] / total_essencial).round(3)
            cesta_pivot = cesta_pivot.sort_values('Total_Acumulado', ascending=False).reset_index()

            # Escrever tabela de subgrupos
            cesta_start_r = 5
            ws_cesta.cell(row=cesta_start_r, column=2, value="Subgrupo Cesta Básica")
            cols = list(cesta_pivot.columns)
            for c_idx, c_name in enumerate(cols):
                ws_cesta.cell(row=cesta_start_r, column=c_idx+2, value=str(c_name))
            self._apply_header_style(ws_cesta, min_row=cesta_start_r, min_col=2, max_col=len(cols)+1, title_bg="385623")

            curr_r = cesta_start_r + 1
            for _, r in cesta_pivot.iterrows():
                for c_idx, c_name in enumerate(cols):
                    val = r[c_name]
                    ws_cesta.cell(row=curr_r, column=c_idx+2, value=val)
                curr_r += 1
            cesta_sub_end_r = curr_r - 1

            # Formatação da tabela de subgrupos
            sub_col_formats = {2: 'left'}
            for c_idx, c_name in enumerate(cols[1:], start=3):
                if c_name == '%_da_Cesta':
                    sub_col_formats[c_idx] = 'percent'
                else:
                    sub_col_formats[c_idx] = 'currency'
            self._format_table(ws_cesta, min_row=cesta_start_r+1, max_row=cesta_sub_end_r, col_formats=sub_col_formats)

            # Gráfico de Subgrupos da Cesta Básica (BarChart horizontal)
            chart_cesta = BarChart()
            chart_cesta.type = "bar"
            chart_cesta.style = 10
            chart_cesta.title = "Gastos por Subgrupo da Cesta Básica (€)"
            chart_cesta.x_axis.title = "Total Gasto (€)"
            chart_cesta.y_axis.title = "Subgrupo"
            chart_cesta.width = 16
            chart_cesta.height = 10
            chart_cesta.legend = None
            tot_col_idx = cols.index('Total_Acumulado') + 2
            data_cesta = Reference(ws_cesta, min_col=tot_col_idx, min_row=cesta_start_r, max_row=cesta_sub_end_r)
            cats_cesta = Reference(ws_cesta, min_col=2, min_row=cesta_start_r+1, max_row=cesta_sub_end_r)
            chart_cesta.add_data(data_cesta, titles_from_data=True)
            chart_cesta.set_categories(cats_cesta)
            ws_cesta.add_chart(chart_cesta, f"B{cesta_sub_end_r + 3}")

            # Top 40 Itens mais consumidos da Cesta Básica
            top_itens_cesta = df_cesta.groupby(['Produto', 'Subgrupo_Cesta']).agg(
                Vezes_Comprado=('Produto', 'count'),
                Qtd_Total=('Quantidade', 'sum'),
                Gasto_Total=('Total', 'sum')
            ).reset_index()
            top_itens_cesta['%_da_Cesta'] = (top_itens_cesta['Gasto_Total'] / total_essencial).round(3)
            top_itens_cesta = top_itens_cesta.sort_values('Gasto_Total', ascending=False).head(40)

            top_start_r = 5
            top_start_col = len(cols) + 4
            ws_cesta.cell(row=top_start_r-1, column=top_start_col, value="Top 40 Produtos Essenciais (Cesta Básica)").font = Font(name="Calibri", size=12, bold=True)
            top_cols = ['Produto', 'Subgrupo', 'Vezes', 'Qtd', 'Total Gasto', '% da Cesta']
            for idx, name in enumerate(top_cols):
                ws_cesta.cell(row=top_start_r, column=top_start_col + idx, value=name)
            self._apply_header_style(ws_cesta, min_row=top_start_r, min_col=top_start_col, max_col=top_start_col+len(top_cols)-1, title_bg="385623")

            curr_r = top_start_r + 1
            for _, r in top_itens_cesta.iterrows():
                ws_cesta.cell(row=curr_r, column=top_start_col, value=str(r['Produto']))
                ws_cesta.cell(row=curr_r, column=top_start_col+1, value=str(r['Subgrupo_Cesta']))
                ws_cesta.cell(row=curr_r, column=top_start_col+2, value=int(r['Vezes_Comprado']))
                ws_cesta.cell(row=curr_r, column=top_start_col+3, value=float(r['Qtd_Total']))
                ws_cesta.cell(row=curr_r, column=top_start_col+4, value=float(r['Gasto_Total']))
                ws_cesta.cell(row=curr_r, column=top_start_col+5, value=float(r['%_da_Cesta']))
                curr_r += 1
            top_end_r = curr_r - 1
            self._format_table(ws_cesta, min_row=top_start_r+1, max_row=top_end_r, col_formats={
                top_start_col: 'left',
                top_start_col+1: 'left',
                top_start_col+2: 'integer',
                top_start_col+3: 'decimal',
                top_start_col+4: 'currency',
                top_start_col+5: 'percent'
            })
            self._auto_adjust_columns(ws_cesta)

            # ==============================================================
            # 3. ABA MENSAL POR CATEGORIA (Visão Geral de Todas as Categorias)
            # ==============================================================
            mensal_cat = df.groupby(['Categoria', 'Mês'])['Total'].sum().reset_index()
            mensal_pivot = mensal_cat.pivot(index='Categoria', columns='Mês', values='Total').fillna(0)
            mensal_pivot['Gasto_Total_Periodo'] = mensal_pivot.sum(axis=1)
            mensal_pivot['%_do_Orçamento'] = (mensal_pivot['Gasto_Total_Periodo'] / total_global).round(3)
            mensal_pivot = mensal_pivot.sort_values('Gasto_Total_Periodo', ascending=False).reset_index()

            ws_mensal = writer.book.create_sheet(title="3. Mensal por Categoria")
            ws_mensal.views.sheetView[0].showGridLines = True
            ws_mensal['B2'] = "GASTOS MENSAIS POR CATEGORIA"
            ws_mensal['B2'].font = Font(name="Calibri", size=15, bold=True, color="1F4E78")

            m_cols = list(mensal_pivot.columns)
            m_start_r = 4
            for c_idx, c_name in enumerate(m_cols):
                ws_mensal.cell(row=m_start_r, column=c_idx+2, value=str(c_name))
            self._apply_header_style(ws_mensal, min_row=m_start_r, min_col=2, max_col=len(m_cols)+1)

            curr_r = m_start_r + 1
            for _, r in mensal_pivot.iterrows():
                for c_idx, c_name in enumerate(m_cols):
                    ws_mensal.cell(row=curr_r, column=c_idx+2, value=r[c_name])
                curr_r += 1
            m_end_r = curr_r - 1

            m_formats = {2: 'left'}
            for c_idx, c_name in enumerate(m_cols[1:], start=3):
                if c_name == '%_do_Orçamento':
                    m_formats[c_idx] = 'percent'
                else:
                    m_formats[c_idx] = 'currency'
            self._format_table(ws_mensal, min_row=m_start_r+1, max_row=m_end_r, col_formats=m_formats)
            self._auto_adjust_columns(ws_mensal)

            # ==============================================================
            # 4. ABA VILÕES DO ORÇAMENTO (Curva ABC / Top Produtos)
            # ==============================================================
            viloes = df.groupby(['Produto', 'Categoria']).agg(
                Vezes_Comprado=('Produto', 'count'),
                Qtd_Total=('Quantidade', 'sum'),
                Gasto_Total=('Total', 'sum')
            ).reset_index()
            viloes['%_do_Orçamento'] = (viloes['Gasto_Total'] / total_global).round(3)
            viloes = viloes.sort_values('Gasto_Total', ascending=False).head(100)

            ws_viloes = writer.book.create_sheet(title="4. Vilões do Orçamento")
            ws_viloes.views.sheetView[0].showGridLines = True
            ws_viloes['B2'] = "TOP 100 PRODUTOS (ONDE VAI O SEU DINHEIRO)"
            ws_viloes['B2'].font = Font(name="Calibri", size=15, bold=True, color="C00000")

            v_cols = ['Produto', 'Categoria', 'Vezes Comprado', 'Qtd Total', 'Gasto Total', '% Orçamento Global']
            v_start_r = 4
            for idx, name in enumerate(v_cols):
                ws_viloes.cell(row=v_start_r, column=idx+2, value=name)
            self._apply_header_style(ws_viloes, min_row=v_start_r, min_col=2, max_col=len(v_cols)+1, title_bg="C00000")

            curr_r = v_start_r + 1
            for _, r in viloes.iterrows():
                ws_viloes.cell(row=curr_r, column=2, value=str(r['Produto']))
                ws_viloes.cell(row=curr_r, column=3, value=str(r['Categoria']))
                ws_viloes.cell(row=curr_r, column=4, value=int(r['Vezes_Comprado']))
                ws_viloes.cell(row=curr_r, column=5, value=float(r['Qtd_Total']))
                ws_viloes.cell(row=curr_r, column=6, value=float(r['Gasto_Total']))
                ws_viloes.cell(row=curr_r, column=7, value=float(r['%_do_Orçamento']))
                curr_r += 1
            v_end_r = curr_r - 1
            self._format_table(ws_viloes, min_row=v_start_r+1, max_row=v_end_r, col_formats={
                2: 'left', 3: 'left', 4: 'integer', 5: 'decimal', 6: 'currency', 7: 'percent'
            })
            self._auto_adjust_columns(ws_viloes)

            # ==============================================================
            # 5. ABA DADOS BRUTOS (Para Auditoria)
            # ==============================================================
            ws_dados = writer.book.create_sheet(title="5. Detalhes Brutos")
            ws_dados.views.sheetView[0].showGridLines = True
            df_bruto = df[['Data', 'Mês', 'Loja', 'Fatura', 'Categoria', 'Subgrupo_Cesta', 'Produto', 'Quantidade', 'Total']].copy()
            df_bruto['Data'] = df_bruto['Data'].dt.strftime('%Y-%m-%d')
            b_cols = list(df_bruto.columns)
            b_start_r = 1
            for c_idx, c_name in enumerate(b_cols, start=1):
                ws_dados.cell(row=b_start_r, column=c_idx, value=str(c_name))
            self._apply_header_style(ws_dados, min_row=b_start_r, min_col=1, max_col=len(b_cols), title_bg="595959")

            curr_r = b_start_r + 1
            for _, r in df_bruto.iterrows():
                for c_idx, c_name in enumerate(b_cols, start=1):
                    ws_dados.cell(row=curr_r, column=c_idx, value=r[c_name])
                curr_r += 1
            b_end_r = curr_r - 1
            tot_idx = b_cols.index('Total') + 1
            qtd_idx = b_cols.index('Quantidade') + 1
            self._format_table(ws_dados, min_row=b_start_r+1, max_row=b_end_r, col_formats={
                tot_idx: 'currency',
                qtd_idx: 'decimal'
            })
            self._auto_adjust_columns(ws_dados)

        # Guarda também uma cópia com o nome padrão relatorio_faturas.xlsx para conveniência
        if output_path != latest_path:
            shutil.copyfile(output_path, latest_path)

        print(f"Relatório gerado com sucesso em:")
        print(f" 📄 Ficheiro desta execução: {output_path}")
        print(f" 📄 Cópia mais recente:      {latest_path}")
        return output_path
